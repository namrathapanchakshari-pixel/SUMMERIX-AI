import os

from docx import Document
from PyPDF2 import PdfReader


def _validate_file(file_path, label):
    if not file_path:
        raise ValueError(f"{label} file path is empty.")

    if not os.path.isfile(file_path):
        raise FileNotFoundError("The uploaded file was not found.")


def extract_pdf_text(file_path):
    _validate_file(file_path, "PDF")

    try:
        reader = PdfReader(file_path)
        pages = []

        for page in reader.pages:
            text = (page.extract_text() or "").strip()

            if text:
                pages.append(text)

        result = "\n\n".join(pages).strip()

        if not result:
            raise ValueError(
                "No selectable text was found in this PDF. "
                "Scanned PDFs need OCR support."
            )

        return result

    except ValueError:
        raise

    except Exception as error:
        raise RuntimeError(
            f"Could not read this PDF: {error}"
        ) from error


def extract_docx_text(file_path):
    _validate_file(file_path, "DOCX")

    try:
        document = Document(file_path)
        content = []

        for paragraph in document.paragraphs:
            text = paragraph.text.strip()

            if text:
                content.append(text)

        for table in document.tables:
            for row in table.rows:
                cells = []

                for cell in row.cells:
                    cell_text = cell.text.strip()

                    if cell_text:
                        cells.append(cell_text)

                if cells:
                    content.append(" | ".join(cells))

        result = "\n\n".join(content).strip()

        if not result:
            raise ValueError(
                "No readable text was found in this DOCX file."
            )

        return result

    except ValueError:
        raise

    except Exception as error:
        raise RuntimeError(
            f"Could not read this DOCX file: {error}"
        ) from error


def extract_text(file_path):
    extension = os.path.splitext(
        str(file_path or "")
    )[1].lower()

    if extension == ".pdf":
        return extract_pdf_text(file_path)

    if extension == ".docx":
        return extract_docx_text(file_path)

    raise ValueError(
        "Only PDF and DOCX files are supported."
    )


def extract_document(file_path):
    """
    Return complete document text and labelled chunks.
    PDF chunks use real page numbers.
    DOCX chunks use numbered sections.
    """
    extension = os.path.splitext(
        str(file_path or "")
    )[1].lower()

    document_name = os.path.basename(file_path)

    if extension == ".pdf":
        _validate_file(file_path, "PDF")

        try:
            reader = PdfReader(file_path)
            chunks = []

            for page_number, page in enumerate(
                reader.pages,
                start=1,
            ):
                text = (page.extract_text() or "").strip()

                if text:
                    chunks.append(
                        {
                            "source": document_name,
                            "location": f"page {page_number}",
                            "text": text,
                        }
                    )

            if not chunks:
                raise ValueError(
                    "No selectable text was found in this PDF. "
                    "Scanned PDFs need OCR support."
                )

            complete_text = "\n\n".join(
                chunk["text"]
                for chunk in chunks
            )

            return {
                "name": document_name,
                "text": complete_text,
                "chunks": chunks,
            }

        except ValueError:
            raise

        except Exception as error:
            raise RuntimeError(
                f"Could not read this PDF: {error}"
            ) from error

    if extension == ".docx":
        complete_text = extract_docx_text(file_path)

        chunk_size = 3500
        chunks = []

        for index, start in enumerate(
            range(0, len(complete_text), chunk_size),
            start=1,
        ):
            chunks.append(
                {
                    "source": document_name,
                    "location": f"section {index}",
                    "text": complete_text[
                        start:start + chunk_size
                    ],
                }
            )

        return {
            "name": document_name,
            "text": complete_text,
            "chunks": chunks,
        }

    raise ValueError(
        "Only PDF and DOCX files are supported."
    )