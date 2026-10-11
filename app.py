import io
import os
import re
from functools import wraps

import sqlite3

import uuid

from datetime import datetime

from dotenv import load_dotenv

from flask import (

    Flask,

    jsonify,

    redirect,

    render_template,

    request,

    send_file,

    session,

    url_for,

)

from werkzeug.security import generate_password_hash, check_password_hash
from werkzeug.utils import secure_filename
from authlib.integrations.flask_client import OAuth

from ai_summarizer import (

    answer_document_question,

    answer_result_question,

    process_document_tool,

)

from chatbot_ai import (

    chat_reply,

    search_videos,

    study_reply,

)

from document_processor import (

    extract_document,

    extract_text,

)

BASE_DIR = os.path.abspath(os.path.dirname(__file__))

load_dotenv(

    os.path.join(BASE_DIR, ".env"),

    override=True,

)

UPLOAD_FOLDER = os.path.join(BASE_DIR, "uploads")

STATE_FOLDER = os.path.join(BASE_DIR, "runtime_state")

DATABASE_PATH = os.path.join(STATE_FOLDER, "summarix.db")

ALLOWED_EXTENSIONS = {".pdf", ".docx"}

RESULTS = {}

DOCUMENTS = {}

os.makedirs(UPLOAD_FOLDER, exist_ok=True)

os.makedirs(STATE_FOLDER, exist_ok=True)

app = Flask(__name__)

app.secret_key = os.getenv(

    "FLASK_SECRET_KEY",

    "change-this-secret-key",

)

app.config["MAX_CONTENT_LENGTH"] = 25 * 1024 * 1024

# Google Sign-In (configure GOOGLE_CLIENT_ID and GOOGLE_CLIENT_SECRET in .env).
oauth = OAuth(app)
google = oauth.register(
    name="google",
    client_id=os.getenv("GOOGLE_CLIENT_ID"),
    client_secret=os.getenv("GOOGLE_CLIENT_SECRET"),
    server_metadata_url="https://accounts.google.com/.well-known/openid-configuration",
    client_kwargs={"scope": "openid email profile"},
)

def google_oauth_is_configured():
    return bool(os.getenv("GOOGLE_CLIENT_ID") and os.getenv("GOOGLE_CLIENT_SECRET"))

def get_database():

    connection = sqlite3.connect(DATABASE_PATH)

    connection.row_factory = sqlite3.Row

    return connection

def initialize_database():

    with get_database() as connection:

        connection.executescript(

            """

            CREATE TABLE IF NOT EXISTS conversations (

                id TEXT PRIMARY KEY,

                user_id INTEGER,

                title TEXT NOT NULL,

                created_at TEXT NOT NULL

            );

            CREATE TABLE IF NOT EXISTS messages (

                id INTEGER PRIMARY KEY AUTOINCREMENT,

                conversation_id TEXT NOT NULL,

                role TEXT NOT NULL,

                content TEXT NOT NULL,

                created_at TEXT NOT NULL

            );

            CREATE TABLE IF NOT EXISTS admins (

                id INTEGER PRIMARY KEY AUTOINCREMENT,

                username TEXT UNIQUE NOT NULL,

                password TEXT NOT NULL,

                created_at TEXT NOT NULL

            );

            CREATE TABLE IF NOT EXISTS users (

                id INTEGER PRIMARY KEY AUTOINCREMENT,

                username TEXT UNIQUE NOT NULL,

                email TEXT UNIQUE,

                created_at TEXT NOT NULL

            );

            CREATE TABLE IF NOT EXISTS documents (

                id INTEGER PRIMARY KEY AUTOINCREMENT,

                user_id INTEGER,

                filename TEXT NOT NULL,

                uploaded_at TEXT NOT NULL

            );

            CREATE TABLE IF NOT EXISTS activity_logs (

                id INTEGER PRIMARY KEY AUTOINCREMENT,

                user_id INTEGER,

                action TEXT NOT NULL,

                details TEXT,

                created_at TEXT NOT NULL

            );

            """

        )

        # Add password_hash to older databases without deleting existing users.
        user_columns = {row[1] for row in connection.execute("PRAGMA table_info(users)")}
        if "password_hash" not in user_columns:
            connection.execute("ALTER TABLE users ADD COLUMN password_hash TEXT")

        conversation_columns = {row[1] for row in connection.execute("PRAGMA table_info(conversations)")}
        if "user_id" not in conversation_columns:
            connection.execute("ALTER TABLE conversations ADD COLUMN user_id INTEGER")


initialize_database()

def allowed_file(filename):

    extension = os.path.splitext(filename)[1].lower()

    return extension in ALLOWED_EXTENSIONS

def safe_upload_path(file_path):

    try:

        return (

            os.path.commonpath(

                [

                    os.path.abspath(file_path),

                    os.path.abspath(UPLOAD_FOLDER),

                ]

            )

            == os.path.abspath(UPLOAD_FOLDER)

        )

    except ValueError:

        return False

def current_document():

    file_path = session.get("document_path")

    document_name = session.get("document_name")

    if (

        not file_path

        or not document_name

        or not safe_upload_path(file_path)

        or not os.path.exists(file_path)

    ):

        raise ValueError(

            "Please upload a PDF or DOCX document first."

        )

    text = extract_text(file_path)

    if not text.strip():

        raise ValueError(

            "No readable text was found in the document."

        )

    return document_name, text

def selected_chat_documents():

    document_ids = session.get("chat_document_ids", [])

    return [

        DOCUMENTS[document_id]

        for document_id in document_ids

        if document_id in DOCUMENTS

    ]

def save_message(conversation_id, role, content):

    with get_database() as connection:

        connection.execute(

            """

            INSERT INTO messages (

                conversation_id,

                role,

                content,

                created_at

            )

            VALUES (?, ?, ?, ?)

            """,

            (

                conversation_id,

                role,

                content,

                datetime.utcnow().isoformat(),

            ),

        )

def ensure_conversation(conversation_id, first_text="New conversation"):
    conversation_id = str(conversation_id or uuid.uuid4().hex)
    user_id = session["user_id"]
    with get_database() as connection:
        existing = connection.execute(
            "SELECT id FROM conversations WHERE id = ? AND user_id = ?",
            (conversation_id, user_id),
        ).fetchone()
        if not existing:
            # If a client submits an ID belonging to another account, use a new ID.
            if connection.execute("SELECT 1 FROM conversations WHERE id = ?", (conversation_id,)).fetchone():
                conversation_id = uuid.uuid4().hex
            title = " ".join(str(first_text).split())[:55] or "New conversation"
            connection.execute(
                "INSERT INTO conversations (id, user_id, title, created_at) VALUES (?, ?, ?, ?)",
                (conversation_id, user_id, title, datetime.utcnow().isoformat()),
            )
    return conversation_id

def login_required(view):
    @wraps(view)
    def wrapped_view(*args, **kwargs):
        if "user_id" not in session:
            if request.path.startswith("/api/"):
                return jsonify(success=False, error="Please log in to continue."), 401
            return redirect(url_for("user_login"))
        return view(*args, **kwargs)
    return wrapped_view


@app.route("/")
def home():
    if "user_id" in session:
        return redirect(url_for("workspace"))
    if "admin_id" in session:
        return redirect(url_for("admin_dashboard"))
    return redirect(url_for("user_login"))

@app.route("/workspace", methods=["GET", "POST"])
@login_required
def workspace():

    result = None

    error = None

    formatting = {}

    active_tool = request.args.get(

        "tool",

        "summary",

    )

    custom_query = ""

    if request.method == "POST":

        active_tool = request.form.get(

            "tool",

            "summary",

        )

        custom_query = request.form.get(

            "custom_query",

            "",

        ).strip()

        try:

            uploaded_file = request.files.get("document")

            if uploaded_file and uploaded_file.filename:

                filename = secure_filename(

                    uploaded_file.filename

                )

                if not allowed_file(filename):

                    raise ValueError(

                        "Only PDF and DOCX files are supported."

                    )

                file_path = os.path.join(

                    UPLOAD_FOLDER,

                    f"{uuid.uuid4().hex}_{filename}",

                )

                uploaded_file.save(file_path)

                with get_database() as connection:
                    connection.execute(
                        "INSERT INTO documents (user_id, filename, uploaded_at) VALUES (?, ?, ?)",
                        (session["user_id"], filename, datetime.now().isoformat(timespec="seconds")),
                    )
                    connection.execute(
                        "INSERT INTO activity_logs (user_id, action, details, created_at) VALUES (?, ?, ?, ?)",
                        (session["user_id"], "Document uploaded", filename, datetime.now().isoformat(timespec="seconds")),
                    )

                session["document_path"] = file_path

                session["document_name"] = filename

            _, document_text = current_document()

            result_data = process_document_tool(

                text=document_text,

                tool=active_tool,

                custom_query=custom_query,

                summary_length=request.form.get(

                    "summary_length",

                    "medium",

                ),

                summary_style=request.form.get(

                    "summary_style",

                    "paragraph",

                ),

                paraphrase_style=request.form.get(

                    "paraphrase_style",

                    "simple",

                ),

                question_count=request.form.get(

                    "question_count",

                    10,

                ),

            )

            result = result_data["content"]

            formatting = result_data["formatting"]

            result_id = uuid.uuid4().hex

            RESULTS[result_id] = {

                "result": result,

                "document": document_text,

            }

            session["last_result_id"] = result_id

        except Exception as exception:

            error = str(exception)

    return render_template(

        "workspace.html",

        result=result,

        error=error,

        formatting=formatting,

        document_name=session.get("document_name"),

        active_tool=active_tool,

        custom_query=custom_query,

        result_id=session.get("last_result_id", ""),

    )

@app.route("/api/document-question", methods=["POST"])
@login_required
def document_question():

    try:

        data = request.get_json(silent=True) or {}

        _, document_text = current_document()

        answer = answer_document_question(

            document_text,

            data.get("question", ""),

        )

        return jsonify(

            success=True,

            answer=answer,

        )

    except Exception as exception:

        return jsonify(

            success=False,

            error=str(exception),

        ), 400

@app.route("/api/result-question", methods=["POST"])
@login_required
def result_question():

    try:

        data = request.get_json(silent=True) or {}

        result_id = (

            data.get("result_id")

            or session.get("last_result_id")

        )

        item = RESULTS.get(result_id)

        if not item:

            raise ValueError(

                "Generate a result before asking about it."

            )

        answer = answer_result_question(

            item["result"],

            data.get("question", ""),

            item["document"],

        )

        return jsonify(

            success=True,

            answer=answer,

        )

    except Exception as exception:

        return jsonify(

            success=False,

            error=str(exception),

        ), 400

@app.route(

    "/api/chat/documents",

    methods=["GET", "POST", "DELETE"],

)

@login_required
def chat_documents():

    if request.method == "DELETE":

        session["chat_document_ids"] = []

        return jsonify(success=True)

    if request.method == "POST":

        try:

            uploaded_files = request.files.getlist(

                "documents"

            )

            if not uploaded_files:

                raise ValueError(

                    "Choose one or more PDF or DOCX files."

                )

            document_ids = session.get(

                "chat_document_ids",

                [],

            )

            for uploaded_file in uploaded_files[:5]:

                filename = secure_filename(

                    uploaded_file.filename

                )

                if (

                    not filename

                    or not allowed_file(filename)

                ):

                    raise ValueError(

                        "Only PDF and DOCX files are supported."

                    )

                file_path = os.path.join(

                    UPLOAD_FOLDER,

                    f"{uuid.uuid4().hex}_{filename}",

                )

                uploaded_file.save(file_path)

                with get_database() as connection:
                    connection.execute(
                        "INSERT INTO documents (user_id, filename, uploaded_at) VALUES (?, ?, ?)",
                        (session["user_id"], filename, datetime.now().isoformat(timespec="seconds")),
                    )
                    connection.execute(
                        "INSERT INTO activity_logs (user_id, action, details, created_at) VALUES (?, ?, ?, ?)",
                        (session["user_id"], "Chat document uploaded", filename, datetime.now().isoformat(timespec="seconds")),
                    )

                document_id = uuid.uuid4().hex

                extracted_document = extract_document(

                    file_path

                )

                DOCUMENTS[document_id] = {

                    "id": document_id,

                    "path": file_path,

                    **extracted_document,

                }

                document_ids.append(document_id)

            session["chat_document_ids"] = document_ids[-5:]

            return jsonify(

                success=True,

                documents=[

                    {

                        "id": item["id"],

                        "name": item["name"],

                    }

                    for item in selected_chat_documents()

                ],

            )

        except Exception as exception:

            return jsonify(

                success=False,

                error=str(exception),

            ), 400

    return jsonify(

        success=True,

        documents=[

            {

                "id": item["id"],

                "name": item["name"],

            }

            for item in selected_chat_documents()

        ],

    )

@app.route(

    "/api/chat/documents/<document_id>",

    methods=["DELETE"],

)

@login_required
def remove_chat_document(document_id):

    session["chat_document_ids"] = [

        item

        for item in session.get(

            "chat_document_ids",

            [],

        )

        if item != document_id

    ]

    return jsonify(success=True)

@app.route("/api/chat", methods=["POST"])
@login_required
def chat():

    try:

        data = request.get_json(silent=True) or {}

        messages = data.get("messages", [])

        if not isinstance(messages, list):

            raise ValueError("Invalid chat history.")

        user_text = next(

            (

                item.get("content", "")

                for item in reversed(messages)

                if item.get("role") == "user"

            ),

            "",

        )

        conversation_id = ensure_conversation(

            data.get("conversation_id"),

            user_text,

        )

        answer, sources = chat_reply(

            messages=messages,

            mode=data.get("mode", "general"),

            documents=selected_chat_documents(),

            preferences=data.get("preferences", {}),

        )

        save_message(

            conversation_id,

            "user",

            user_text,

        )

        save_message(

            conversation_id,

            "assistant",

            answer,

        )

        return jsonify(

            success=True,

            conversation_id=conversation_id,

            answer=answer,

            sources=sources,

        )

    except Exception as exception:

        return jsonify(

            success=False,

            error=str(exception),

        ), 400

@app.route("/api/study", methods=["POST"])
@login_required
def study():

    try:

        data = request.get_json(silent=True) or {}

        answer, sources = study_reply(

            action=data.get("action"),

            topic=data.get("topic"),

            documents=selected_chat_documents(),

            preferences=data.get("preferences", {}),

        )

        return jsonify(

            success=True,

            answer=answer,

            sources=sources,

        )

    except Exception as exception:

        return jsonify(

            success=False,

            error=str(exception),

        ), 400

@app.route("/api/chat/videos", methods=["POST"])
@login_required
def videos():

    data = request.get_json(silent=True) or {}

    return jsonify(

        success=True,

        videos=search_videos(

            data.get("query", "")

        ),

    )

@app.route("/api/conversations", methods=["GET", "DELETE"])
@login_required
def conversations():
    user_id = session["user_id"]
    with get_database() as connection:
        if request.method == "DELETE":
            connection.execute(
                "DELETE FROM messages WHERE conversation_id IN (SELECT id FROM conversations WHERE user_id = ?)",
                (user_id,),
            )
            connection.execute("DELETE FROM conversations WHERE user_id = ?", (user_id,))
            return jsonify(success=True)
        rows = connection.execute(
            "SELECT id, title, created_at FROM conversations WHERE user_id = ? ORDER BY created_at DESC LIMIT 30",
            (user_id,),
        ).fetchall()
    return jsonify(success=True, conversations=[dict(row) for row in rows])

@app.route("/api/conversations/<conversation_id>", methods=["GET", "DELETE"])
@login_required
def conversation(conversation_id):
    user_id = session["user_id"]
    with get_database() as connection:
        owned = connection.execute(
            "SELECT id FROM conversations WHERE id = ? AND user_id = ?",
            (conversation_id, user_id),
        ).fetchone()
        if not owned:
            return jsonify(success=False, error="Conversation not found."), 404
        if request.method == "DELETE":
            connection.execute("DELETE FROM messages WHERE conversation_id = ?", (conversation_id,))
            connection.execute("DELETE FROM conversations WHERE id = ? AND user_id = ?", (conversation_id, user_id))
            return jsonify(success=True)
        rows = connection.execute(
            "SELECT role, content FROM messages WHERE conversation_id = ? ORDER BY id",
            (conversation_id,),
        ).fetchall()
    return jsonify(success=True, messages=[dict(row) for row in rows])

@app.route("/api/chat/export", methods=["POST"])
@login_required
def export_chat():

    data = request.get_json(silent=True) or {}

    messages = data.get("messages", [])[-100:]

    export_format = data.get("format", "txt")

    content = "SUMMARIX AI CHAT EXPORT\n\n"

    for item in messages:

        role = str(

            item.get("role", "assistant")

        ).title()

        message = str(item.get("content", ""))

        content += f"{role}: {message}\n\n"

    if export_format == "docx":

        from docx import Document

        document = Document()

        document.add_heading(

            "SUMMARIX AI Chat Export",

            0,

        )

        for item in messages:

            document.add_heading(

                str(

                    item.get(

                        "role",

                        "assistant",

                    )

                ).title(),

                2,

            )

            document.add_paragraph(

                str(item.get("content", ""))

            )

        output = io.BytesIO()

        document.save(output)

        output.seek(0)

        return send_file(

            output,

            as_attachment=True,

            download_name="summarix-chat.docx",

            mimetype=(

                "application/vnd.openxmlformats-officedocument."

                "wordprocessingml.document"

            ),

        )

    if export_format == "pdf":

        from reportlab.lib.pagesizes import A4

        from reportlab.pdfgen import canvas

        output = io.BytesIO()

        pdf = canvas.Canvas(

            output,

            pagesize=A4,

        )

        text = pdf.beginText(40, 800)

        for line in content.splitlines():

            parts = [

                line[index:index + 95]

                for index in range(

                    0,

                    max(1, len(line)),

                    95,

                )

            ]

            for part in parts:

                text.textLine(part)

                if text.getY() < 45:

                    pdf.drawText(text)

                    pdf.showPage()

                    text = pdf.beginText(40, 800)

        pdf.drawText(text)

        pdf.save()

        output.seek(0)

        return send_file(

            output,

            as_attachment=True,

            download_name="summarix-chat.pdf",

            mimetype="application/pdf",

        )

    output = io.BytesIO(

        content.encode("utf-8")

    )

    return send_file(

        output,

        as_attachment=True,

        download_name="summarix-chat.txt",

        mimetype="text/plain",

    )

@app.errorhandler(413)

def file_too_large(error):

    return jsonify(

        success=False,

        error="File is too large. Maximum size is 25 MB.",

    ), 413

@app.route("/register", methods=["GET", "POST"])
def user_register():
    error = None
    if request.method == "POST":
        username = request.form.get("username", "").strip()
        email = request.form.get("email", "").strip().lower()
        password = request.form.get("password", "")
        confirm_password = request.form.get("confirm_password", "")

        if len(username) < 3 or len(username) > 40:
            error = "Username must be between 3 and 40 characters."
        elif not email or "@" not in email:
            error = "Enter a valid email address."
        elif len(password) < 8:
            error = "Password must contain at least 8 characters."
        elif password != confirm_password:
            error = "Passwords do not match."
        else:
            try:
                with get_database() as connection:
                    cursor = connection.execute(
                        "INSERT INTO users (username, email, created_at, password_hash) VALUES (?, ?, ?, ?)",
                        (username, email, datetime.now().isoformat(timespec="seconds"), generate_password_hash(password)),
                    )
                    user_id = cursor.lastrowid
                session.clear()
                session["user_id"] = user_id
                session["username"] = username
                return redirect(url_for("workspace"))
            except sqlite3.IntegrityError:
                error = "That username or email is already registered."

    return render_template("register.html", error=error, google_enabled=google_oauth_is_configured())


@app.route("/login", methods=["GET", "POST"])
def user_login():
    error = None
    if request.method == "POST":
        username_or_email = request.form.get("username_or_email", "").strip()
        password = request.form.get("password", "")
        with get_database() as connection:
            user = connection.execute(
                "SELECT id, username, email, password_hash FROM users WHERE username = ? OR email = ?",
                (username_or_email, username_or_email.lower()),
            ).fetchone()

        if user and user["password_hash"] and check_password_hash(user["password_hash"], password):
            session.clear()
            session["user_id"] = user["id"]
            session["username"] = user["username"]
            return redirect(url_for("workspace"))
        error = "Invalid username/email or password."

    return render_template("login.html", error=error, google_enabled=google_oauth_is_configured())


@app.route("/auth/google")
def google_login():
    if not google_oauth_is_configured():
        return render_template(
            "login.html",
            error="Google Sign-In is not configured yet. Add GOOGLE_CLIENT_ID and GOOGLE_CLIENT_SECRET to your .env file.",
            google_enabled=False,
        ), 503
    callback_url = url_for("google_callback", _external=True)
    return google.authorize_redirect(callback_url)


@app.route("/auth/google/callback")
def google_callback():
    if not google_oauth_is_configured():
        return redirect(url_for("user_login"))

    try:
        token = google.authorize_access_token()
        profile = token.get("userinfo") or {}
        email = (profile.get("email") or "").strip().lower()
        email_verified = profile.get("email_verified", False)
        display_name = (profile.get("name") or profile.get("given_name") or "Google User").strip()

        if not email or email_verified is not True:
            return render_template(
                "login.html",
                error="Google did not return a verified email address. Please try another account.",
                google_enabled=True,
            ), 400

        with get_database() as connection:
            user = connection.execute(
                "SELECT id, username, email FROM users WHERE email = ?", (email,)
            ).fetchone()

            if user:
                user_id = user["id"]
                username = user["username"]
            else:
                base = re.sub(r"[^a-zA-Z0-9_.-]", "", display_name.replace(" ", ""))[:30]
                if len(base) < 3:
                    base = email.split("@", 1)[0]
                base = re.sub(r"[^a-zA-Z0-9_.-]", "", base)[:30] or "googleuser"
                username = base
                suffix = 1
                while connection.execute(
                    "SELECT 1 FROM users WHERE username = ?", (username,)
                ).fetchone():
                    username = f"{base[:34]}{suffix}"
                    suffix += 1
                cursor = connection.execute(
                    "INSERT INTO users (username, email, created_at, password_hash) VALUES (?, ?, ?, NULL)",
                    (username, email, datetime.now().isoformat(timespec="seconds")),
                )
                user_id = cursor.lastrowid

        session.clear()
        session["user_id"] = user_id
        session["username"] = username
        session["email"] = email
        return redirect(url_for("workspace"))
    except Exception:
        app.logger.exception("Google sign-in failed")
        return render_template(
            "login.html",
            error="Google sign-in could not be completed. Check your Google OAuth settings and try again.",
            google_enabled=google_oauth_is_configured(),
        ), 400


@app.route("/profile", methods=["GET", "POST"])
@login_required
def user_profile():
    message = None
    error = None
    user_id = session["user_id"]
    if request.method == "POST":
        username = request.form.get("username", "").strip()
        email = request.form.get("email", "").strip().lower()
        if len(username) < 3 or len(username) > 40:
            error = "Username must be between 3 and 40 characters."
        elif not email or not re.match(r"^[^\s@]+@[^\s@]+\.[^\s@]+$", email):
            error = "Enter a valid email address."
        else:
            try:
                with get_database() as connection:
                    connection.execute(
                        "UPDATE users SET username = ?, email = ? WHERE id = ?",
                        (username, email, user_id),
                    )
                session["username"] = username
                session["email"] = email
                message = "Profile updated successfully."
            except sqlite3.IntegrityError:
                error = "That username or email is already in use."
    with get_database() as connection:
        user = connection.execute(
            "SELECT id, username, email, created_at FROM users WHERE id = ?", (user_id,)
        ).fetchone()
    if not user:
        session.clear()
        return redirect(url_for("user_login"))
    return render_template("profile.html", user=user, message=message, error=error)

@app.route("/settings", methods=["GET", "POST"])
@login_required
def user_settings():
    message = None
    error = None
    if request.method == "POST":
        current_password = request.form.get("current_password", "")
        new_password = request.form.get("new_password", "")
        confirm_password = request.form.get("confirm_password", "")
        if len(new_password) < 8:
            error = "New password must contain at least 8 characters."
        elif new_password != confirm_password:
            error = "New passwords do not match."
        else:
            with get_database() as connection:
                user = connection.execute(
                    "SELECT password_hash FROM users WHERE id = ?", (session["user_id"],)
                ).fetchone()
                if not user:
                    error = "Account not found. Please sign in again."
                elif not user["password_hash"]:
                    error = "This account uses Google Sign-In and does not have a local password to change."
                elif not check_password_hash(user["password_hash"], current_password):
                    error = "Current password is incorrect."
                else:
                    connection.execute(
                        "UPDATE users SET password_hash = ? WHERE id = ?",
                        (generate_password_hash(new_password), session["user_id"]),
                    )
                    message = "Password changed successfully."
    return render_template("settings.html", message=message, error=error)

@app.route("/history")
@login_required
def user_history():
    with get_database() as connection:
        documents = connection.execute(
            "SELECT id, filename, uploaded_at FROM documents WHERE user_id = ? ORDER BY id DESC LIMIT 200",
            (session["user_id"],),
        ).fetchall()
    return render_template("history.html", documents=documents)

@app.route("/logout")
def user_logout():
    session.clear()
    return redirect(url_for("user_login"))


@app.route("/admin/login", methods=["GET", "POST"])

def admin_login():

    error = None

    if request.method == "POST":

        username = request.form.get("username", "").strip()

        password = request.form.get("password", "")

        with get_database() as connection:

            admin = connection.execute(

                """

                SELECT id, username, password

                FROM admins

                WHERE username = ?

                """,

                (username,),

            ).fetchone()

        if admin and check_password_hash(admin["password"], password):

            session["admin_id"] = admin["id"]

            session["admin_username"] = admin["username"]

            return redirect(url_for("admin_dashboard"))

        error = "Invalid username or password."

    return render_template("admin_login.html", error=error)

@app.route("/admin/dashboard")

def admin_dashboard():

    if "admin_id" not in session:

        return redirect(url_for("admin_login"))

    with get_database() as connection:

        admin_count = connection.execute(

            "SELECT COUNT(*) FROM admins"

        ).fetchone()[0]

        user_count = connection.execute(

            "SELECT COUNT(*) FROM users"

        ).fetchone()[0]

        document_count = connection.execute(

            "SELECT COUNT(*) FROM documents"

        ).fetchone()[0]

        activity_count = connection.execute(

            "SELECT COUNT(*) FROM activity_logs"

        ).fetchone()[0]

        recent_activities = connection.execute(

            """

            SELECT action, details, created_at

            FROM activity_logs

            ORDER BY id DESC

            LIMIT 10

            """

        ).fetchall()

    return render_template(

        "admin_dashboard.html",

        admin_username=session.get("admin_username", "Admin"),

        admin_count=admin_count,

        user_count=user_count,

        document_count=document_count,

        activity_count=activity_count,

        recent_activities=recent_activities,

    )

@app.route("/admin/logout")

def admin_logout():

    session.pop("admin_id", None)

    session.pop("admin_username", None)

    return redirect(url_for("admin_login"))

if __name__ == "__main__":

    print(

        "SUMMARIX AI: "

        "http://127.0.0.1:5000/workspace"

    )

    app.run(

        host="127.0.0.1",

        port=5000,

        debug=True,

    )
